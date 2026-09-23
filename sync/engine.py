from dataclasses import dataclass, field

import openpyxl
import psycopg

from sync.workbook_reader import read_list_values, read_tab_rows


@dataclass
class SyncReport:
    rows_written: dict = field(default_factory=dict)
    errors: list = field(default_factory=list)


def sync_workbook_to_supabase(filepath: str, farm_code: str, dsn: str) -> SyncReport:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    return _sync_from_workbook(wb, farm_code, dsn)


def _sync_from_workbook(wb, farm_code: str, dsn: str) -> SyncReport:
    """The actual sync logic, decoupled from *how* `wb` was loaded - an
    openpyxl Workbook (sync_workbook_to_supabase, above) or a
    sync.sheets_io.SheetsWorkbook (sync_sheet_to_supabase, below) both
    satisfy the same minimal interface (`.sheetnames`, `wb[name]` returning
    something read_tab_rows/read_list_values can read), so every _sync_*
    function below runs unchanged against either source."""
    # prepare_threshold=None disables psycopg3's server-side prepared
    # statements - required against Supabase's transaction-pooler DSN
    # (pgbouncer transaction mode), which can route each statement to a
    # different backend connection; a statement prepared on one backend
    # doesn't exist on the next, surfacing as a real
    # 'prepared statement "_pg3_N" does not exist' error mid-sync.
    conn = psycopg.connect(dsn, autocommit=True, prepare_threshold=None)
    report = SyncReport()

    lists = read_list_values(wb["99_LISTS"]) if "99_LISTS" in wb.sheetnames else {}

    report.rows_written["farm_profile"] = _sync_farm_profile(wb, conn)
    report.rows_written["staff"] = _sync_staff(wb, conn, farm_code, lists, report.errors)
    if lists:
        report.rows_written["list_values"] = _sync_list_values(conn, lists)
    if "P1_PIG_BATCHES" in wb.sheetnames:
        report.rows_written["pig_batches"] = _sync_pig_batches(
            wb, conn, farm_code, lists, report.errors)
    if "P2_PIG_DAILY_LOG" in wb.sheetnames:
        report.rows_written["pig_daily_log"] = _sync_pig_daily_log(
            wb, conn, report.errors)
    if "P3_PIG_WEIGHTS" in wb.sheetnames:
        report.rows_written["pig_weights"] = _sync_pig_weights(
            wb, conn, report.errors)
    if "P4_SOW_REGISTER" in wb.sheetnames:
        report.rows_written["sow_register"] = _sync_sow_register(
            wb, conn, farm_code, lists, report.errors)
    if "P5_BREEDING_FARROWING" in wb.sheetnames:
        report.rows_written["breeding_farrowing"] = _sync_breeding_farrowing(
            wb, conn, lists, report.errors)
    if "C1_POULTRY_BATCHES" in wb.sheetnames:
        report.rows_written["poultry_batches"] = _sync_poultry_batches(
            wb, conn, farm_code, lists, report.errors)
    if "C2_POULTRY_DAILY_LOG" in wb.sheetnames:
        report.rows_written["poultry_daily_log"] = _sync_poultry_daily_log(
            wb, conn, report.errors)
    if "C3_POULTRY_WEIGHTS" in wb.sheetnames:
        report.rows_written["poultry_weights"] = _sync_poultry_weights(
            wb, conn, report.errors)
    if "F1_PLOTS" in wb.sheetnames:
        report.rows_written["plots"] = _sync_plots(
            wb, conn, farm_code, lists, report.errors)
    if "F2_PLANTINGS" in wb.sheetnames:
        report.rows_written["plantings"] = _sync_plantings(
            wb, conn, lists, report.errors)
    if "F3_FIELD_OPERATIONS" in wb.sheetnames:
        report.rows_written["field_operations"] = _sync_field_operations(
            wb, conn, lists, report.errors)
    if "F4_SCOUTING_LOG" in wb.sheetnames:
        report.rows_written["scouting_log"] = _sync_scouting_log(
            wb, conn, lists, report.errors)
    if "F5_HARVEST_LOG" in wb.sheetnames:
        report.rows_written["harvest_log"] = _sync_harvest_log(
            wb, conn, lists, report.errors)
    if "06_FEED_INVENTORY" in wb.sheetnames:
        report.rows_written["feed_inventory"] = _sync_feed_inventory(
            wb, conn, lists, report.errors)
    if "05_WEATHER_LOG" in wb.sheetnames:
        report.rows_written["weather_log"] = _sync_weather_log(wb, conn)
    if "04_LABOUR_LOG" in wb.sheetnames:
        report.rows_written["labour_log"] = _sync_labour_log(wb, conn, report.errors)
    if "02_EXPENSES" in wb.sheetnames:
        report.rows_written["expenses"] = _sync_expenses(wb, conn, lists, report.errors)
    if "03_REVENUE" in wb.sheetnames:
        report.rows_written["revenue"] = _sync_revenue(wb, conn, lists, report.errors)
    if "07_HEALTH_LOG" in wb.sheetnames:
        report.rows_written["health_log"] = _sync_health_log(wb, conn, report.errors)

    conn.close()
    return report


def sync_sheet_to_supabase(sheet_id: str, farm_code: str, dsn: str, creds_json: str) -> SyncReport:
    """Same sync, sourced from the live Google Sheet instead of the .xlsx
    workbook. `creds_json` is the service account key's raw JSON content
    (a string), not a file path. See sync/sheets_io.py for how a Sheet is
    made to look like an openpyxl Workbook to the _sync_* functions below."""
    from sync.sheets_io import load_sheets_workbook
    wb = load_sheets_workbook(sheet_id, creds_json)
    return _sync_from_workbook(wb, farm_code, dsn)


PIG_BATCH_DROPDOWN_FIELDS = {
    "breed": "BREED_PIG",
    "source": "SOURCE",
    "status": "STATUS_BATCH",
}

SOW_DROPDOWN_FIELDS = {
    "breed": "BREED_PIG",
    "status": "SOW_STATUS",
}

FARROWING_DROPDOWN_FIELDS = {
    "service_type": "SERVICE_TYPE",
    "pregnancy_confirmed": "YESNO",
}

POULTRY_BATCH_DROPDOWN_FIELDS = {
    "bird_type": "BIRD_TYPE",
    "breed": "BREED_POULTRY",
    "status": "POULTRY_STATUS",
}

PLOT_DROPDOWN_FIELDS = {
    "soil_type": "SOIL_TYPE",
    "irrigation_type": "IRRIGATION_TYPE",
}

PLANTING_DROPDOWN_FIELDS = {
    "crop": "CROP",
    "season": "SEASON",
    "status": "PLANTING_STATUS",
}

FIELD_OPERATION_DROPDOWN_FIELDS = {
    "op_type": "OP_TYPE",
    "unit": "UNIT_FIELD",
}

SCOUTING_DROPDOWN_FIELDS = {
    "growth_stage": "GROWTH_STAGE",
    "issue_type": "ISSUE_TYPE",
    "soil_moisture": "SOIL_MOISTURE",
}

HARVEST_DROPDOWN_FIELDS = {
    "grade": "GRADE",
    "storage_location": "STORAGE",
}


# staff column -> 99_LISTS column it must be one of
STAFF_DROPDOWN_FIELDS = {
    "primary_domain": "DOMAIN",
    "role": "ROLE",
}

REVENUE_DROPDOWN_FIELDS = {
    "payment_status": "PAYMENT_STATUS",
}


def _validate_dropdowns(table: str, record: dict, row_id: str,
                         dropdown_fields: dict, lists: dict, errors: list) -> bool:
    """Returns True if the record passes all dropdown checks; else appends
    one error per violation to `errors` and returns False."""
    valid = True
    for field_name, list_name in dropdown_fields.items():
        if list_name not in lists:
            continue  # no 99_LISTS data available to validate against
        value = record[field_name]
        if value not in lists[list_name]:
            errors.append({
                "table": table, "row": row_id, "field": field_name,
                "value": value, "reason": f"not in 99_LISTS[{list_name}]",
            })
            valid = False
    return valid


BATCH_REF_TABLE_BY_DOMAIN = {
    "piggery": "pig_batches",
    "poultry": "poultry_batches",
}


def _validate_batch_ref(table: str, record: dict, row_id: str,
                         batch_ref_field: str, conn, errors: list) -> bool:
    """For a nullable, polymorphic batch reference: if present, it must
    resolve in the batch table for that row's domain (piggery->pig_batches,
    poultry->poultry_batches). Other domains aren't checked - there's no
    single anchor table for them at this reference's grain."""
    batch_ref = record[batch_ref_field]
    if batch_ref is None:
        return True
    ref_table = BATCH_REF_TABLE_BY_DOMAIN.get(record["domain"])
    if ref_table is None:
        return True
    exists = conn.execute(
        f"SELECT 1 FROM {ref_table} WHERE batch_code = %s", (batch_ref,)
    ).fetchone()
    if exists is None:
        errors.append({
            "table": table, "row": row_id, "field": batch_ref_field,
            "value": batch_ref, "reason": f"no matching batch_code in {ref_table}",
        })
        return False
    return True


def _sync_farm_profile(wb, conn) -> int:
    # Singleton table: read from row 3 (real workbooks have no row 4 at
    # all here), and take the LAST row found so a fixture that models
    # row 3 as a discardable sample plus a real row 4 still resolves to
    # the real one.
    record = read_tab_rows(wb["00_FARM_PROFILE"], min_row=3)[-1]

    conn.execute("""
        CREATE TABLE IF NOT EXISTS farm_profile (
            farm_code TEXT PRIMARY KEY,
            farm_name TEXT,
            region_district TEXT,
            currency TEXT,
            total_hectares NUMERIC,
            module_piggery_active BOOLEAN,
            module_poultry_active BOOLEAN,
            module_crops_active BOOLEAN,
            financial_year_start DATE
        )
    """)
    conn.execute(
        """
        INSERT INTO farm_profile (
            farm_code, farm_name, region_district, currency, total_hectares,
            module_piggery_active, module_poultry_active, module_crops_active,
            financial_year_start
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (farm_code) DO UPDATE SET
            farm_name = EXCLUDED.farm_name,
            region_district = EXCLUDED.region_district,
            currency = EXCLUDED.currency,
            total_hectares = EXCLUDED.total_hectares,
            module_piggery_active = EXCLUDED.module_piggery_active,
            module_poultry_active = EXCLUDED.module_poultry_active,
            module_crops_active = EXCLUDED.module_crops_active,
            financial_year_start = EXCLUDED.financial_year_start
        """,
        (
            record["farm_code"], record["farm_name"], record["region_district"],
            record["currency"], record["total_hectares"],
            record["module_piggery_active"] == "Yes",
            record["module_poultry_active"] == "Yes",
            record["module_crops_active"] == "Yes",
            record["financial_year_start"],
        ),
    )
    return 1


def _sync_staff(wb, conn, farm_code: str, lists: dict, errors: list) -> int:
    # Row 3 holds a real staff member (S01), not a discardable sample -
    # see read_tab_rows' docstring.
    records = read_tab_rows(wb["01_STAFF"], min_row=3)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS staff (
            staff_code TEXT PRIMARY KEY,
            farm_code TEXT NOT NULL REFERENCES farm_profile(farm_code),
            full_name TEXT,
            role TEXT,
            primary_domain TEXT,
            pay_type TEXT,
            rate NUMERIC,
            active BOOLEAN
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("staff", record, record["staff_code"],
                                    STAFF_DROPDOWN_FIELDS, lists, errors):
            continue
        conn.execute(
            """
            INSERT INTO staff (
                staff_code, farm_code, full_name, role, primary_domain, pay_type,
                rate, active
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (staff_code) DO UPDATE SET
                farm_code = EXCLUDED.farm_code,
                full_name = EXCLUDED.full_name,
                role = EXCLUDED.role,
                primary_domain = EXCLUDED.primary_domain,
                pay_type = EXCLUDED.pay_type,
                rate = EXCLUDED.rate,
                active = EXCLUDED.active
            """,
            (
                record["staff_code"], farm_code, record["full_name"], record["role"],
                record["primary_domain"], record["pay_type"], record["rate"],
                record["active"] == "Yes",
            ),
        )
        written += 1
    return written


def _sync_pig_batches(wb, conn, farm_code: str, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["P1_PIG_BATCHES"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS pig_batches (
            batch_code TEXT PRIMARY KEY,
            farm_code TEXT NOT NULL REFERENCES farm_profile(farm_code),
            pen TEXT,
            breed TEXT,
            start_date DATE,
            start_count INTEGER,
            start_avg_kg NUMERIC,
            source TEXT,
            target_market_date DATE,
            status TEXT
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("pig_batches", record, record["batch_code"],
                                    PIG_BATCH_DROPDOWN_FIELDS, lists, errors):
            continue
        conn.execute(
            """
            INSERT INTO pig_batches (
                batch_code, farm_code, pen, breed, start_date, start_count,
                start_avg_kg, source, target_market_date, status
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (batch_code) DO UPDATE SET
                farm_code = EXCLUDED.farm_code,
                pen = EXCLUDED.pen,
                breed = EXCLUDED.breed,
                start_date = EXCLUDED.start_date,
                start_count = EXCLUDED.start_count,
                start_avg_kg = EXCLUDED.start_avg_kg,
                source = EXCLUDED.source,
                target_market_date = EXCLUDED.target_market_date,
                status = EXCLUDED.status
            """,
            (
                record["batch_code"], farm_code, record["pen"], record["breed"],
                record["start_date"], record["start_count"], record["start_avg_kg"],
                record["source"], record["target_market_date"], record["status"],
            ),
        )
        written += 1
    return written


def _sync_pig_daily_log(wb, conn, errors: list) -> int:
    records = read_tab_rows(wb["P2_PIG_DAILY_LOG"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS pig_daily_log (
            date DATE NOT NULL,
            batch_code TEXT NOT NULL REFERENCES pig_batches(batch_code),
            pen TEXT,
            feed_type TEXT,
            feed_kg NUMERIC,
            water_litres NUMERIC,
            deaths INTEGER,
            culls INTEGER,
            closing_count INTEGER,
            pen_temp_c NUMERIC,
            pen_condition TEXT,
            notes TEXT,
            recorded_by TEXT,
            PRIMARY KEY (date, batch_code)
        )
    """)
    written = 0
    for record in records:
        try:
            conn.execute(
                """
                INSERT INTO pig_daily_log (
                    date, batch_code, pen, feed_type, feed_kg, water_litres,
                    deaths, culls, closing_count, pen_temp_c, pen_condition,
                    notes, recorded_by
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, batch_code) DO UPDATE SET
                    pen = EXCLUDED.pen,
                    feed_type = EXCLUDED.feed_type,
                    feed_kg = EXCLUDED.feed_kg,
                    water_litres = EXCLUDED.water_litres,
                    deaths = EXCLUDED.deaths,
                    culls = EXCLUDED.culls,
                    closing_count = EXCLUDED.closing_count,
                    pen_temp_c = EXCLUDED.pen_temp_c,
                    pen_condition = EXCLUDED.pen_condition,
                    notes = EXCLUDED.notes,
                    recorded_by = EXCLUDED.recorded_by
                """,
                (
                    record["date"], record["batch_code"], record["pen"],
                    record["feed_type"], record["feed_kg"], record["water_litres"],
                    record["deaths"], record["culls"], record["closing_count"],
                    record["pen_temp_c"], record["pen_condition"], record["notes"],
                    record["recorded_by"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "pig_daily_log",
                "row": f"{record['date']}/{record['batch_code']}",
                "field": "batch_code", "value": record["batch_code"],
                "reason": "no matching batch_code in pig_batches",
            })
            continue
        written += 1
    return written


def _sync_pig_weights(wb, conn, errors: list) -> int:
    records = read_tab_rows(wb["P3_PIG_WEIGHTS"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS pig_weights (
            date DATE NOT NULL,
            batch_code TEXT NOT NULL REFERENCES pig_batches(batch_code),
            sample_size INTEGER,
            avg_weight_kg NUMERIC,
            min_weight_kg NUMERIC,
            max_weight_kg NUMERIC,
            age_days INTEGER,
            PRIMARY KEY (date, batch_code)
        )
    """)
    written = 0
    for record in records:
        try:
            conn.execute(
                """
                INSERT INTO pig_weights (
                    date, batch_code, sample_size, avg_weight_kg, min_weight_kg,
                    max_weight_kg, age_days
                ) VALUES (%s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, batch_code) DO UPDATE SET
                    sample_size = EXCLUDED.sample_size,
                    avg_weight_kg = EXCLUDED.avg_weight_kg,
                    min_weight_kg = EXCLUDED.min_weight_kg,
                    max_weight_kg = EXCLUDED.max_weight_kg,
                    age_days = EXCLUDED.age_days
                """,
                (
                    record["date"], record["batch_code"], record["sample_size"],
                    record["avg_weight_kg"], record["min_weight_kg"],
                    record["max_weight_kg"], record["age_days"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "pig_weights",
                "row": f"{record['date']}/{record['batch_code']}",
                "field": "batch_code", "value": record["batch_code"],
                "reason": "no matching batch_code in pig_batches",
            })
            continue
        written += 1
    return written


def _sync_sow_register(wb, conn, farm_code: str, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["P4_SOW_REGISTER"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS sow_register (
            sow_tag TEXT PRIMARY KEY,
            farm_code TEXT NOT NULL REFERENCES farm_profile(farm_code),
            breed TEXT,
            date_of_birth DATE,
            parity INTEGER,
            body_condition_1_5 INTEGER,
            status TEXT
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("sow_register", record, record["sow_tag"],
                                    SOW_DROPDOWN_FIELDS, lists, errors):
            continue
        conn.execute(
            """
            INSERT INTO sow_register (
                sow_tag, farm_code, breed, date_of_birth, parity,
                body_condition_1_5, status
            ) VALUES (%s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (sow_tag) DO UPDATE SET
                farm_code = EXCLUDED.farm_code,
                breed = EXCLUDED.breed,
                date_of_birth = EXCLUDED.date_of_birth,
                parity = EXCLUDED.parity,
                body_condition_1_5 = EXCLUDED.body_condition_1_5,
                status = EXCLUDED.status
            """,
            (
                record["sow_tag"], farm_code, record["breed"],
                record["date_of_birth"], record["parity"],
                record["body_condition_1_5"], record["status"],
            ),
        )
        written += 1
    return written


def _sync_breeding_farrowing(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["P5_BREEDING_FARROWING"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS breeding_farrowing (
            sow_tag TEXT NOT NULL REFERENCES sow_register(sow_tag),
            boar_tag TEXT,
            service_date DATE NOT NULL,
            service_type TEXT,
            pregnancy_confirmed BOOLEAN,
            expected_farrow_date DATE,
            farrow_date DATE,
            born_alive INTEGER,
            stillborn INTEGER,
            mummified INTEGER,
            avg_birth_weight_kg NUMERIC,
            wean_date DATE,
            weaned_count INTEGER,
            avg_wean_weight_kg NUMERIC,
            litter_batch_code TEXT REFERENCES pig_batches(batch_code),
            PRIMARY KEY (sow_tag, service_date)
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("breeding_farrowing", record, record["sow_tag"],
                                    FARROWING_DROPDOWN_FIELDS, lists, errors):
            continue
        try:
            conn.execute(
                """
                INSERT INTO breeding_farrowing (
                    sow_tag, boar_tag, service_date, service_type,
                    pregnancy_confirmed, expected_farrow_date, farrow_date,
                    born_alive, stillborn, mummified, avg_birth_weight_kg,
                    wean_date, weaned_count, avg_wean_weight_kg, litter_batch_code
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (sow_tag, service_date) DO UPDATE SET
                    boar_tag = EXCLUDED.boar_tag,
                    service_type = EXCLUDED.service_type,
                    pregnancy_confirmed = EXCLUDED.pregnancy_confirmed,
                    expected_farrow_date = EXCLUDED.expected_farrow_date,
                    farrow_date = EXCLUDED.farrow_date,
                    born_alive = EXCLUDED.born_alive,
                    stillborn = EXCLUDED.stillborn,
                    mummified = EXCLUDED.mummified,
                    avg_birth_weight_kg = EXCLUDED.avg_birth_weight_kg,
                    wean_date = EXCLUDED.wean_date,
                    weaned_count = EXCLUDED.weaned_count,
                    avg_wean_weight_kg = EXCLUDED.avg_wean_weight_kg,
                    litter_batch_code = EXCLUDED.litter_batch_code
                """,
                (
                    record["sow_tag"], record["boar_tag"], record["service_date"],
                    record["service_type"], record["pregnancy_confirmed"] == "Yes",
                    record["expected_farrow_date"], record["farrow_date"],
                    record["born_alive"], record["stillborn"], record["mummified"],
                    record["avg_birth_weight_kg"], record["wean_date"],
                    record["weaned_count"], record["avg_wean_weight_kg"],
                    record["litter_batch_code"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "breeding_farrowing",
                "row": f"{record['sow_tag']}/{record['service_date']}",
                "field": "sow_tag_or_litter_batch_code",
                "value": (record["sow_tag"], record["litter_batch_code"]),
                "reason": "no matching sow_tag in sow_register or "
                          "litter_batch_code in pig_batches",
            })
            continue
        written += 1
    return written


def _sync_poultry_batches(wb, conn, farm_code: str, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["C1_POULTRY_BATCHES"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS poultry_batches (
            batch_code TEXT PRIMARY KEY,
            farm_code TEXT NOT NULL REFERENCES farm_profile(farm_code),
            bird_type TEXT,
            breed TEXT,
            house TEXT,
            placement_date DATE,
            chicks_placed INTEGER,
            chick_unit_cost NUMERIC,
            hatchery TEXT,
            target_off_date DATE,
            status TEXT
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("poultry_batches", record, record["batch_code"],
                                    POULTRY_BATCH_DROPDOWN_FIELDS, lists, errors):
            continue
        conn.execute(
            """
            INSERT INTO poultry_batches (
                batch_code, farm_code, bird_type, breed, house, placement_date,
                chicks_placed, chick_unit_cost, hatchery, target_off_date, status
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (batch_code) DO UPDATE SET
                farm_code = EXCLUDED.farm_code,
                bird_type = EXCLUDED.bird_type,
                breed = EXCLUDED.breed,
                house = EXCLUDED.house,
                placement_date = EXCLUDED.placement_date,
                chicks_placed = EXCLUDED.chicks_placed,
                chick_unit_cost = EXCLUDED.chick_unit_cost,
                hatchery = EXCLUDED.hatchery,
                target_off_date = EXCLUDED.target_off_date,
                status = EXCLUDED.status
            """,
            (
                record["batch_code"], farm_code, record["bird_type"], record["breed"],
                record["house"], record["placement_date"], record["chicks_placed"],
                record["chick_unit_cost"], record["hatchery"],
                record["target_off_date"], record["status"],
            ),
        )
        written += 1
    return written


def _sync_poultry_daily_log(wb, conn, errors: list) -> int:
    records = read_tab_rows(wb["C2_POULTRY_DAILY_LOG"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS poultry_daily_log (
            date DATE NOT NULL,
            batch_code TEXT NOT NULL REFERENCES poultry_batches(batch_code),
            age_days INTEGER,
            deaths INTEGER,
            culls INTEGER,
            closing_birds INTEGER,
            feed_type TEXT,
            feed_kg NUMERIC,
            water_litres NUMERIC,
            house_temp_c NUMERIC,
            litter_condition TEXT,
            lighting_hours NUMERIC,
            eggs_collected INTEGER,
            eggs_cracked INTEGER,
            eggs_dirty INTEGER,
            notes TEXT,
            recorded_by TEXT,
            PRIMARY KEY (date, batch_code)
        )
    """)
    written = 0
    for record in records:
        try:
            conn.execute(
                """
                INSERT INTO poultry_daily_log (
                    date, batch_code, age_days, deaths, culls, closing_birds,
                    feed_type, feed_kg, water_litres, house_temp_c,
                    litter_condition, lighting_hours, eggs_collected,
                    eggs_cracked, eggs_dirty, notes, recorded_by
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, batch_code) DO UPDATE SET
                    age_days = EXCLUDED.age_days,
                    deaths = EXCLUDED.deaths,
                    culls = EXCLUDED.culls,
                    closing_birds = EXCLUDED.closing_birds,
                    feed_type = EXCLUDED.feed_type,
                    feed_kg = EXCLUDED.feed_kg,
                    water_litres = EXCLUDED.water_litres,
                    house_temp_c = EXCLUDED.house_temp_c,
                    litter_condition = EXCLUDED.litter_condition,
                    lighting_hours = EXCLUDED.lighting_hours,
                    eggs_collected = EXCLUDED.eggs_collected,
                    eggs_cracked = EXCLUDED.eggs_cracked,
                    eggs_dirty = EXCLUDED.eggs_dirty,
                    notes = EXCLUDED.notes,
                    recorded_by = EXCLUDED.recorded_by
                """,
                (
                    record["date"], record["batch_code"], record["age_days"],
                    record["deaths"], record["culls"], record["closing_birds"],
                    record["feed_type"], record["feed_kg"], record["water_litres"],
                    record["house_temp_c"], record["litter_condition"],
                    record["lighting_hours"], record["eggs_collected"],
                    record["eggs_cracked"], record["eggs_dirty"], record["notes"],
                    record["recorded_by"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "poultry_daily_log",
                "row": f"{record['date']}/{record['batch_code']}",
                "field": "batch_code", "value": record["batch_code"],
                "reason": "no matching batch_code in poultry_batches",
            })
            continue
        written += 1
    return written


def _sync_poultry_weights(wb, conn, errors: list) -> int:
    records = read_tab_rows(wb["C3_POULTRY_WEIGHTS"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS poultry_weights (
            date DATE NOT NULL,
            batch_code TEXT NOT NULL REFERENCES poultry_batches(batch_code),
            age_days INTEGER,
            sample_size INTEGER,
            avg_weight_g NUMERIC,
            uniformity_pct NUMERIC,
            PRIMARY KEY (date, batch_code)
        )
    """)
    written = 0
    for record in records:
        try:
            conn.execute(
                """
                INSERT INTO poultry_weights (
                    date, batch_code, age_days, sample_size, avg_weight_g,
                    uniformity_pct
                ) VALUES (%s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, batch_code) DO UPDATE SET
                    age_days = EXCLUDED.age_days,
                    sample_size = EXCLUDED.sample_size,
                    avg_weight_g = EXCLUDED.avg_weight_g,
                    uniformity_pct = EXCLUDED.uniformity_pct
                """,
                (
                    record["date"], record["batch_code"], record["age_days"],
                    record["sample_size"], record["avg_weight_g"],
                    record["uniformity_pct"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "poultry_weights",
                "row": f"{record['date']}/{record['batch_code']}",
                "field": "batch_code", "value": record["batch_code"],
                "reason": "no matching batch_code in poultry_batches",
            })
            continue
        written += 1
    return written


def _sync_plots(wb, conn, farm_code: str, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["F1_PLOTS"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS plots (
            plot_code TEXT PRIMARY KEY,
            farm_code TEXT NOT NULL REFERENCES farm_profile(farm_code),
            plot_name TEXT,
            hectares NUMERIC,
            soil_type TEXT,
            irrigation_type TEXT,
            gps TEXT
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("plots", record, record["plot_code"],
                                    PLOT_DROPDOWN_FIELDS, lists, errors):
            continue
        conn.execute(
            """
            INSERT INTO plots (
                plot_code, farm_code, plot_name, hectares, soil_type,
                irrigation_type, gps
            ) VALUES (%s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (plot_code) DO UPDATE SET
                farm_code = EXCLUDED.farm_code,
                plot_name = EXCLUDED.plot_name,
                hectares = EXCLUDED.hectares,
                soil_type = EXCLUDED.soil_type,
                irrigation_type = EXCLUDED.irrigation_type,
                gps = EXCLUDED.gps
            """,
            (
                record["plot_code"], farm_code, record["plot_name"],
                record["hectares"], record["soil_type"],
                record["irrigation_type"], record["gps"],
            ),
        )
        written += 1
    return written


def _sync_plantings(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["F2_PLANTINGS"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS plantings (
            planting_code TEXT PRIMARY KEY,
            plot_code TEXT NOT NULL REFERENCES plots(plot_code),
            crop TEXT,
            variety TEXT,
            season TEXT,
            planting_date DATE,
            area_ha NUMERIC,
            seed_kg NUMERIC,
            seed_cost NUMERIC,
            expected_harvest_date DATE,
            expected_yield_tons_ha NUMERIC,
            status TEXT
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("plantings", record, record["planting_code"],
                                    PLANTING_DROPDOWN_FIELDS, lists, errors):
            continue
        try:
            conn.execute(
                """
                INSERT INTO plantings (
                    planting_code, plot_code, crop, variety, season,
                    planting_date, area_ha, seed_kg, seed_cost,
                    expected_harvest_date, expected_yield_tons_ha, status
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (planting_code) DO UPDATE SET
                    plot_code = EXCLUDED.plot_code,
                    crop = EXCLUDED.crop,
                    variety = EXCLUDED.variety,
                    season = EXCLUDED.season,
                    planting_date = EXCLUDED.planting_date,
                    area_ha = EXCLUDED.area_ha,
                    seed_kg = EXCLUDED.seed_kg,
                    seed_cost = EXCLUDED.seed_cost,
                    expected_harvest_date = EXCLUDED.expected_harvest_date,
                    expected_yield_tons_ha = EXCLUDED.expected_yield_tons_ha,
                    status = EXCLUDED.status
                """,
                (
                    record["planting_code"], record["plot_code"], record["crop"],
                    record["variety"], record["season"], record["planting_date"],
                    record["area_ha"], record["seed_kg"], record["seed_cost"],
                    record["expected_harvest_date"],
                    record["expected_yield_tons_ha"], record["status"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "plantings", "row": record["planting_code"],
                "field": "plot_code", "value": record["plot_code"],
                "reason": "no matching plot_code in plots",
            })
            continue
        written += 1
    return written


def _sync_field_operations(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["F3_FIELD_OPERATIONS"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS field_operations (
            date DATE NOT NULL,
            planting_code TEXT NOT NULL REFERENCES plantings(planting_code),
            op_type TEXT NOT NULL,
            product TEXT,
            rate TEXT,
            quantity NUMERIC,
            unit TEXT,
            input_cost NUMERIC,
            labour_hours NUMERIC,
            machine_hours NUMERIC,
            irrigation_mm NUMERIC,
            notes TEXT,
            recorded_by TEXT,
            PRIMARY KEY (date, planting_code, op_type)
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("field_operations", record, record["planting_code"],
                                    FIELD_OPERATION_DROPDOWN_FIELDS, lists, errors):
            continue
        try:
            conn.execute(
                """
                INSERT INTO field_operations (
                    date, planting_code, op_type, product, rate, quantity, unit,
                    input_cost, labour_hours, machine_hours, irrigation_mm,
                    notes, recorded_by
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, planting_code, op_type) DO UPDATE SET
                    product = EXCLUDED.product,
                    rate = EXCLUDED.rate,
                    quantity = EXCLUDED.quantity,
                    unit = EXCLUDED.unit,
                    input_cost = EXCLUDED.input_cost,
                    labour_hours = EXCLUDED.labour_hours,
                    machine_hours = EXCLUDED.machine_hours,
                    irrigation_mm = EXCLUDED.irrigation_mm,
                    notes = EXCLUDED.notes,
                    recorded_by = EXCLUDED.recorded_by
                """,
                (
                    record["date"], record["planting_code"], record["op_type"],
                    record["product"], record["rate"], record["quantity"],
                    record["unit"], record["input_cost"], record["labour_hours"],
                    record["machine_hours"], record["irrigation_mm"],
                    record["notes"], record["recorded_by"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "field_operations",
                "row": f"{record['date']}/{record['planting_code']}",
                "field": "planting_code", "value": record["planting_code"],
                "reason": "no matching planting_code in plantings",
            })
            continue
        written += 1
    return written


def _sync_scouting_log(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["F4_SCOUTING_LOG"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS scouting_log (
            date DATE NOT NULL,
            planting_code TEXT NOT NULL REFERENCES plantings(planting_code),
            growth_stage TEXT,
            issue_type TEXT,
            issue_name TEXT,
            severity_1_5 INTEGER,
            incidence_pct NUMERIC,
            soil_moisture TEXT,
            action_taken TEXT,
            scouted_by TEXT,
            PRIMARY KEY (date, planting_code)
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("scouting_log", record, record["planting_code"],
                                    SCOUTING_DROPDOWN_FIELDS, lists, errors):
            continue
        try:
            conn.execute(
                """
                INSERT INTO scouting_log (
                    date, planting_code, growth_stage, issue_type, issue_name,
                    severity_1_5, incidence_pct, soil_moisture, action_taken,
                    scouted_by
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, planting_code) DO UPDATE SET
                    growth_stage = EXCLUDED.growth_stage,
                    issue_type = EXCLUDED.issue_type,
                    issue_name = EXCLUDED.issue_name,
                    severity_1_5 = EXCLUDED.severity_1_5,
                    incidence_pct = EXCLUDED.incidence_pct,
                    soil_moisture = EXCLUDED.soil_moisture,
                    action_taken = EXCLUDED.action_taken,
                    scouted_by = EXCLUDED.scouted_by
                """,
                (
                    record["date"], record["planting_code"], record["growth_stage"],
                    record["issue_type"], record["issue_name"],
                    record["severity_1_5"], record["incidence_pct"],
                    record["soil_moisture"], record["action_taken"],
                    record["scouted_by"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "scouting_log",
                "row": f"{record['date']}/{record['planting_code']}",
                "field": "planting_code", "value": record["planting_code"],
                "reason": "no matching planting_code in plantings",
            })
            continue
        written += 1
    return written


def _sync_harvest_log(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["F5_HARVEST_LOG"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS harvest_log (
            date DATE NOT NULL,
            planting_code TEXT NOT NULL REFERENCES plantings(planting_code),
            bags INTEGER,
            quantity_kg NUMERIC,
            moisture_pct NUMERIC,
            grade TEXT,
            field_loss_kg NUMERIC,
            storage_location TEXT,
            labour_hours NUMERIC,
            PRIMARY KEY (date, planting_code)
        )
    """)
    written = 0
    for record in records:
        if not _validate_dropdowns("harvest_log", record, record["planting_code"],
                                    HARVEST_DROPDOWN_FIELDS, lists, errors):
            continue
        try:
            conn.execute(
                """
                INSERT INTO harvest_log (
                    date, planting_code, bags, quantity_kg, moisture_pct, grade,
                    field_loss_kg, storage_location, labour_hours
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, planting_code) DO UPDATE SET
                    bags = EXCLUDED.bags,
                    quantity_kg = EXCLUDED.quantity_kg,
                    moisture_pct = EXCLUDED.moisture_pct,
                    grade = EXCLUDED.grade,
                    field_loss_kg = EXCLUDED.field_loss_kg,
                    storage_location = EXCLUDED.storage_location,
                    labour_hours = EXCLUDED.labour_hours
                """,
                (
                    record["date"], record["planting_code"], record["bags"],
                    record["quantity_kg"], record["moisture_pct"],
                    record["grade"], record["field_loss_kg"],
                    record["storage_location"], record["labour_hours"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "harvest_log",
                "row": f"{record['date']}/{record['planting_code']}",
                "field": "planting_code", "value": record["planting_code"],
                "reason": "no matching planting_code in plantings",
            })
            continue
        written += 1
    return written


def _sync_feed_inventory(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["06_FEED_INVENTORY"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS feed_inventory (
            date DATE NOT NULL,
            domain TEXT NOT NULL,
            feed_type TEXT NOT NULL,
            opening_kg NUMERIC,
            received_kg NUMERIC,
            used_kg NUMERIC,
            waste_kg NUMERIC,
            closing_kg NUMERIC,
            unit_cost_per_kg NUMERIC,
            supplier TEXT,
            batch_ref TEXT,
            PRIMARY KEY (date, domain, feed_type)
        )
    """)
    written = 0
    for record in records:
        conn.execute(
            """
            INSERT INTO feed_inventory (
                date, domain, feed_type, opening_kg, received_kg, used_kg,
                waste_kg, closing_kg, unit_cost_per_kg, supplier, batch_ref
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (date, domain, feed_type) DO UPDATE SET
                opening_kg = EXCLUDED.opening_kg,
                received_kg = EXCLUDED.received_kg,
                used_kg = EXCLUDED.used_kg,
                waste_kg = EXCLUDED.waste_kg,
                closing_kg = EXCLUDED.closing_kg,
                unit_cost_per_kg = EXCLUDED.unit_cost_per_kg,
                supplier = EXCLUDED.supplier,
                batch_ref = EXCLUDED.batch_ref
            """,
            (
                record["date"], record["domain"], record["feed_type"],
                record["opening_kg"], record["received_kg"], record["used_kg"],
                record["waste_kg"], record["closing_kg"],
                record["unit_cost_per_kg"], record["supplier"],
                record["batch_ref"],
            ),
        )
        written += 1
    return written


def _sync_weather_log(wb, conn) -> int:
    records = read_tab_rows(wb["05_WEATHER_LOG"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS weather_log (
            date DATE PRIMARY KEY,
            rainfall_mm NUMERIC,
            temp_min_c NUMERIC,
            temp_max_c NUMERIC,
            humidity_pct NUMERIC,
            event TEXT
        )
    """)
    for record in records:
        conn.execute(
            """
            INSERT INTO weather_log (
                date, rainfall_mm, temp_min_c, temp_max_c, humidity_pct, event
            ) VALUES (%s, %s, %s, %s, %s, %s)
            ON CONFLICT (date) DO UPDATE SET
                rainfall_mm = EXCLUDED.rainfall_mm,
                temp_min_c = EXCLUDED.temp_min_c,
                temp_max_c = EXCLUDED.temp_max_c,
                humidity_pct = EXCLUDED.humidity_pct,
                event = EXCLUDED.event
            """,
            (
                record["date"], record["rainfall_mm"], record["temp_min_c"],
                record["temp_max_c"], record["humidity_pct"], record["event"],
            ),
        )
    return len(records)


def _sync_labour_log(wb, conn, errors: list) -> int:
    records = read_tab_rows(wb["04_LABOUR_LOG"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS labour_log (
            date DATE NOT NULL,
            staff_code TEXT NOT NULL REFERENCES staff(staff_code),
            domain TEXT,
            task TEXT,
            hours NUMERIC,
            overtime_hours NUMERIC,
            labour_cost NUMERIC,
            batch_ref TEXT,
            PRIMARY KEY (date, staff_code)
        )
    """)
    written = 0
    for record in records:
        row_id = f"{record['date']}/{record['staff_code']}"
        if not _validate_batch_ref("labour_log", record, row_id, "batch_ref",
                                    conn, errors):
            continue
        try:
            conn.execute(
                """
                INSERT INTO labour_log (
                    date, staff_code, domain, task, hours, overtime_hours,
                    labour_cost, batch_ref
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (date, staff_code) DO UPDATE SET
                    domain = EXCLUDED.domain,
                    task = EXCLUDED.task,
                    hours = EXCLUDED.hours,
                    overtime_hours = EXCLUDED.overtime_hours,
                    labour_cost = EXCLUDED.labour_cost,
                    batch_ref = EXCLUDED.batch_ref
                """,
                (
                    record["date"], record["staff_code"], record["domain"],
                    record["task"], record["hours"], record["overtime_hours"],
                    record["labour_cost"], record["batch_ref"],
                ),
            )
        except psycopg.errors.ForeignKeyViolation:
            conn.rollback()
            errors.append({
                "table": "labour_log",
                "row": f"{record['date']}/{record['staff_code']}",
                "field": "staff_code", "value": record["staff_code"],
                "reason": "no matching staff_code in staff",
            })
            continue
        written += 1
    return written


def _sync_expenses(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["02_EXPENSES"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS expenses (
            date DATE NOT NULL,
            domain TEXT,
            category TEXT,
            item TEXT NOT NULL,
            quantity NUMERIC,
            unit TEXT,
            unit_cost NUMERIC,
            total_cost NUMERIC,
            supplier TEXT,
            payment_method TEXT,
            batch_ref TEXT,
            allocation_note TEXT,
            recorded_by TEXT,
            PRIMARY KEY (date, domain, item)
        )
    """)
    written = 0
    for record in records:
        row_id = f"{record['date']}/{record['domain']}/{record['item']}"
        if not _validate_batch_ref("expenses", record, row_id, "batch_ref",
                                    conn, errors):
            continue
        conn.execute(
            """
            INSERT INTO expenses (
                date, domain, category, item, quantity, unit, unit_cost,
                total_cost, supplier, payment_method, batch_ref,
                allocation_note, recorded_by
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (date, domain, item) DO UPDATE SET
                category = EXCLUDED.category,
                quantity = EXCLUDED.quantity,
                unit = EXCLUDED.unit,
                unit_cost = EXCLUDED.unit_cost,
                total_cost = EXCLUDED.total_cost,
                supplier = EXCLUDED.supplier,
                payment_method = EXCLUDED.payment_method,
                batch_ref = EXCLUDED.batch_ref,
                allocation_note = EXCLUDED.allocation_note,
                recorded_by = EXCLUDED.recorded_by
            """,
            (
                record["date"], record["domain"], record["category"],
                record["item"], record["quantity"], record["unit"],
                record["unit_cost"], record["total_cost"], record["supplier"],
                record["payment_method"], record["batch_ref"],
                record["allocation_note"], record["recorded_by"],
            ),
        )
        written += 1
    return written


def _sync_revenue(wb, conn, lists: dict, errors: list) -> int:
    records = read_tab_rows(wb["03_REVENUE"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS revenue (
            date DATE NOT NULL,
            domain TEXT,
            product TEXT,
            quantity NUMERIC,
            unit TEXT,
            unit_price NUMERIC,
            total_amount NUMERIC,
            head_count INTEGER,
            avg_weight_kg NUMERIC,
            grade TEXT,
            buyer TEXT NOT NULL,
            channel TEXT,
            payment_status TEXT,
            batch_ref TEXT,
            PRIMARY KEY (date, domain, buyer)
        )
    """)
    written = 0
    for record in records:
        row_id = f"{record['date']}/{record['domain']}/{record['buyer']}"
        if not _validate_dropdowns("revenue", record, row_id,
                                    REVENUE_DROPDOWN_FIELDS, lists, errors):
            continue
        if not _validate_batch_ref("revenue", record, row_id, "batch_ref",
                                    conn, errors):
            continue
        conn.execute(
            """
            INSERT INTO revenue (
                date, domain, product, quantity, unit, unit_price, total_amount,
                head_count, avg_weight_kg, grade, buyer, channel,
                payment_status, batch_ref
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (date, domain, buyer) DO UPDATE SET
                product = EXCLUDED.product,
                quantity = EXCLUDED.quantity,
                unit = EXCLUDED.unit,
                unit_price = EXCLUDED.unit_price,
                total_amount = EXCLUDED.total_amount,
                head_count = EXCLUDED.head_count,
                avg_weight_kg = EXCLUDED.avg_weight_kg,
                grade = EXCLUDED.grade,
                channel = EXCLUDED.channel,
                payment_status = EXCLUDED.payment_status,
                batch_ref = EXCLUDED.batch_ref
            """,
            (
                record["date"], record["domain"], record["product"],
                record["quantity"], record["unit"], record["unit_price"],
                record["total_amount"], record["head_count"],
                record["avg_weight_kg"], record["grade"], record["buyer"],
                record["channel"], record["payment_status"],
                record["batch_ref"],
            ),
        )
        written += 1
    return written


def _sync_health_log(wb, conn, errors: list) -> int:
    records = read_tab_rows(wb["07_HEALTH_LOG"])

    conn.execute("""
        CREATE TABLE IF NOT EXISTS health_log (
            date DATE NOT NULL,
            domain TEXT,
            batch_ref TEXT NOT NULL,
            animal_tag TEXT,
            event_type TEXT NOT NULL,
            symptom TEXT,
            diagnosis TEXT,
            product TEXT,
            dose TEXT,
            animals_treated INTEGER,
            cost NUMERIC,
            withdrawal_until DATE,
            outcome TEXT,
            administered_by TEXT,
            PRIMARY KEY (date, batch_ref, event_type)
        )
    """)
    written = 0
    for record in records:
        row_id = f"{record['date']}/{record['batch_ref']}/{record['event_type']}"
        if not _validate_batch_ref("health_log", record, row_id, "batch_ref",
                                    conn, errors):
            continue
        conn.execute(
            """
            INSERT INTO health_log (
                date, domain, batch_ref, animal_tag, event_type, symptom,
                diagnosis, product, dose, animals_treated, cost,
                withdrawal_until, outcome, administered_by
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (date, batch_ref, event_type) DO UPDATE SET
                domain = EXCLUDED.domain,
                animal_tag = EXCLUDED.animal_tag,
                symptom = EXCLUDED.symptom,
                diagnosis = EXCLUDED.diagnosis,
                product = EXCLUDED.product,
                dose = EXCLUDED.dose,
                animals_treated = EXCLUDED.animals_treated,
                cost = EXCLUDED.cost,
                withdrawal_until = EXCLUDED.withdrawal_until,
                outcome = EXCLUDED.outcome,
                administered_by = EXCLUDED.administered_by
            """,
            (
                record["date"], record["domain"], record["batch_ref"],
                record["animal_tag"], record["event_type"], record["symptom"],
                record["diagnosis"], record["product"], record["dose"],
                record["animals_treated"], record["cost"],
                record["withdrawal_until"], record["outcome"],
                record["administered_by"],
            ),
        )
        written += 1
    return written


def _sync_list_values(conn, lists: dict) -> int:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS list_values (
            list_name TEXT NOT NULL,
            value TEXT NOT NULL,
            PRIMARY KEY (list_name, value)
        )
    """)
    count = 0
    for list_name, values in lists.items():
        for value in values:
            conn.execute(
                """
                INSERT INTO list_values (list_name, value) VALUES (%s, %s)
                ON CONFLICT (list_name, value) DO NOTHING
                """,
                (list_name, value),
            )
            count += 1
    return count
