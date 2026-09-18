"""
PROTOTYPE-adjacent, but a real fix: replaces the 7 formula-bearing
columns' dead formula strings with their actual computed values.

Discovered: openpyxl cannot evaluate formulas - it only reads whatever
cached value a real spreadsheet engine last saved. This workbook's
cached values were never actually populated (confirmed by direct
inspection: 100% of cells in all 7 formula columns read as None via
data_only=True), contradicting an earlier, apparently unverified claim
that a LibreOffice recalculation had been run against this file.

Rather than depend on installing a full spreadsheet engine, this
script replicates each formula's exact logic in Python (verified by
reading the literal formula strings from the live file first) and
writes the resulting number/date as a static value, replacing the
formula. Scope: only these 7 columns' cells change; nothing else in
the workbook is touched.
"""

import datetime as dt

import openpyxl

WORKBOOK = "../Netrisyl_Farm_Intelligence_Workbook.xlsx"


def read_rows(ws, ncols):
    header = [c.value for c in ws[2]]
    rows = []
    for r in range(4, ws.max_row + 1):
        values = [ws.cell(row=r, column=c + 1).value for c in range(ncols)]
        if any(v is not None for v in values):
            rows.append((r, dict(zip(header, values))))
    return rows


def build_lookup(ws, ncols, key_col, value_col, min_row=4):
    header = [c.value for c in ws[2]]
    key_idx, value_idx = header.index(key_col), header.index(value_col)
    lookup = {}
    for r in range(min_row, ws.max_row + 1):
        key = ws.cell(row=r, column=key_idx + 1).value
        if key is not None:
            lookup[key] = ws.cell(row=r, column=value_idx + 1).value
    return lookup


def fix_expenses(wb):
    ws = wb["02_EXPENSES"]
    header = [c.value for c in ws[2]]
    q_idx, c_idx, t_idx = header.index("quantity"), header.index("unit_cost"), header.index("total_cost")
    n = 0
    for r, record in read_rows(ws, len(header)):
        if record["quantity"] is not None and record["unit_cost"] is not None:
            ws.cell(row=r, column=t_idx + 1, value=round(record["quantity"] * record["unit_cost"], 2))
            n += 1
    return n


def fix_revenue(wb):
    ws = wb["03_REVENUE"]
    header = [c.value for c in ws[2]]
    q_idx, p_idx, t_idx = header.index("quantity"), header.index("unit_price"), header.index("total_amount")
    n = 0
    for r, record in read_rows(ws, len(header)):
        if record["quantity"] is not None and record["unit_price"] is not None:
            ws.cell(row=r, column=t_idx + 1, value=round(record["quantity"] * record["unit_price"], 2))
            n += 1
    return n


def fix_labour_log(wb):
    # 01_STAFF's row 3 holds a real staff member (S01), not a discardable
    # sample - see sync/workbook_reader.py's read_tab_rows docstring for
    # the same finding, discovered independently there.
    staff_rate = build_lookup(wb["01_STAFF"], 7, "staff_code", "rate", min_row=3)
    ws = wb["04_LABOUR_LOG"]
    header = [c.value for c in ws[2]]
    l_idx = header.index("labour_cost")
    n = 0
    for r, record in read_rows(ws, len(header)):
        rate = staff_rate.get(record["staff_code"])
        value = round(rate * record["hours"], 2) if rate is not None and record["hours"] is not None else 0
        ws.cell(row=r, column=l_idx + 1, value=value)
        n += 1
    return n


def fix_pig_daily_log(wb):
    start_count = build_lookup(wb["P1_PIG_BATCHES"], 9, "batch_code", "start_count")
    ws = wb["P2_PIG_DAILY_LOG"]
    header = [c.value for c in ws[2]]
    c_idx = header.index("closing_count")
    n = 0
    for r, record in read_rows(ws, len(header)):
        base = start_count.get(record["batch_code"])
        if base is not None:
            value = base - (record["deaths"] or 0) - (record["culls"] or 0)
        else:
            value = 0
        ws.cell(row=r, column=c_idx + 1, value=value)
        n += 1
    return n


def fix_pig_weights(wb):
    start_date = build_lookup(wb["P1_PIG_BATCHES"], 9, "batch_code", "start_date")
    ws = wb["P3_PIG_WEIGHTS"]
    header = [c.value for c in ws[2]]
    a_idx = header.index("age_days")
    n = 0
    for r, record in read_rows(ws, len(header)):
        base = start_date.get(record["batch_code"])
        if base is not None and record["date"] is not None:
            value = (dt.date.fromisoformat(record["date"]) - dt.date.fromisoformat(base)).days
        else:
            value = ""
        ws.cell(row=r, column=a_idx + 1, value=value)
        n += 1
    return n


def fix_breeding_farrowing(wb):
    ws = wb["P5_BREEDING_FARROWING"]
    header = [c.value for c in ws[2]]
    e_idx = header.index("expected_farrow_date")
    n = 0
    for r, record in read_rows(ws, len(header)):
        if record["service_date"] is not None:
            new_date = dt.date.fromisoformat(record["service_date"]) + dt.timedelta(days=114)
            ws.cell(row=r, column=e_idx + 1, value=new_date.isoformat())
            n += 1
    return n


def fix_harvest_log(wb):
    ws = wb["F5_HARVEST_LOG"]
    header = [c.value for c in ws[2]]
    q_idx = header.index("quantity_kg")
    n = 0
    for r, record in read_rows(ws, len(header)):
        if record["bags"] is not None:
            ws.cell(row=r, column=q_idx + 1, value=record["bags"] * 50)
            n += 1
    return n


if __name__ == "__main__":
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)

    counts = {
        "02_EXPENSES.total_cost": fix_expenses(wb),
        "03_REVENUE.total_amount": fix_revenue(wb),
        "04_LABOUR_LOG.labour_cost": fix_labour_log(wb),
        "P2_PIG_DAILY_LOG.closing_count": fix_pig_daily_log(wb),
        "P3_PIG_WEIGHTS.age_days": fix_pig_weights(wb),
        "P5_BREEDING_FARROWING.expected_farrow_date": fix_breeding_farrowing(wb),
        "F5_HARVEST_LOG.quantity_kg": fix_harvest_log(wb),
    }
    for name, n in counts.items():
        print(f"{name}: {n} cells written")

    wb.save(WORKBOOK)
    print("Saved.")
