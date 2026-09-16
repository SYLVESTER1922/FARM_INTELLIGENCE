import openpyxl
from collections import defaultdict

PATH = "/Users/vharsh/Farm-Intelligence-App/Netrisyl_Farm_Intelligence_Workbook_generated.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=False)

errors = []

def sheet_rows(name, start=4):
    ws = wb[name]
    for r in range(start, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in vals):
            continue
        yield r, vals

# ---------------------------------------------------------------------
# 1. Formula-input sanity: confirm no None feeds into a formula's operands
# ---------------------------------------------------------------------
def check_numeric_formula(sheet, op_cols, label):
    n = 0
    for r, vals in sheet_rows(sheet):
        for c in op_cols:
            if vals[c - 1] is None:
                errors.append(f"{sheet}!row{r}: operand col{c} is None for {label}")
        n += 1
    print(f"{sheet}: checked {n} rows for {label}")

check_numeric_formula("02_EXPENSES", [5, 7], "total_cost = E*G")
check_numeric_formula("03_REVENUE", [4, 6], "total_amount = D*F")
check_numeric_formula("F5_HARVEST_LOG", [3], "quantity_kg = C*50")

# ---------------------------------------------------------------------
# 2. INDEX/MATCH lookups resolve (batch_ref / staff_code exist in anchor sheet)
# ---------------------------------------------------------------------
staff_codes = set()
for r, vals in sheet_rows("01_STAFF", start=3):
    staff_codes.add(vals[0])
print("Known staff codes:", sorted(staff_codes))

missing_staff = set()
for r, vals in sheet_rows("04_LABOUR_LOG"):
    sc = vals[1]
    if sc not in staff_codes:
        missing_staff.add(sc)
        errors.append(f"04_LABOUR_LOG!row{r}: staff_code {sc} not found in 01_STAFF")
print("04_LABOUR_LOG unresolved staff_codes:", missing_staff or "none")

pig_batches = set()
for r, vals in sheet_rows("P1_PIG_BATCHES", start=3):
    pig_batches.add(vals[0])
print("Known pig batch codes:", sorted(pig_batches))

missing_p2 = set()
for r, vals in sheet_rows("P2_PIG_DAILY_LOG"):
    bc = vals[1]
    if bc not in pig_batches:
        missing_p2.add(bc)
        errors.append(f"P2_PIG_DAILY_LOG!row{r}: batch_code {bc} not found in P1_PIG_BATCHES")
print("P2 unresolved batch_codes:", missing_p2 or "none")

missing_p3 = set()
for r, vals in sheet_rows("P3_PIG_WEIGHTS"):
    bc = vals[1]
    if bc not in pig_batches:
        missing_p3.add(bc)
        errors.append(f"P3_PIG_WEIGHTS!row{r}: batch_code {bc} not found in P1_PIG_BATCHES")
print("P3 unresolved batch_codes:", missing_p3 or "none")

poultry_batches = set()
for r, vals in sheet_rows("C1_POULTRY_BATCHES", start=3):
    poultry_batches.add(vals[0])
print("Known poultry batch codes:", sorted(poultry_batches))

missing_c2 = set()
for r, vals in sheet_rows("C2_POULTRY_DAILY_LOG"):
    bc = vals[1]
    if bc not in poultry_batches:
        missing_c2.add(bc)
print("C2 unresolved batch_codes:", missing_c2 or "none")

plantings = set()
for r, vals in sheet_rows("F2_PLANTINGS", start=3):
    plantings.add(vals[0])
print("Known planting codes:", sorted(plantings))

for sheet, col in [("F3_FIELD_OPERATIONS", 2), ("F4_SCOUTING_LOG", 2), ("F5_HARVEST_LOG", 2)]:
    missing = set()
    for r, vals in sheet_rows(sheet):
        pc = vals[col - 1]
        if pc not in plantings:
            missing.add(pc)
            errors.append(f"{sheet}!row{r}: planting_code {pc} not found in F2_PLANTINGS")
    print(f"{sheet} unresolved planting_codes:", missing or "none")

# 02_EXPENSES / 03_REVENUE batch_ref should resolve to a known batch/planting when set
all_batch_like = pig_batches | poultry_batches | plantings
for sheet, col in [("02_EXPENSES", 11), ("03_REVENUE", 14)]:
    missing = set()
    for r, vals in sheet_rows(sheet):
        ref = vals[col - 1]
        if ref is not None and ref not in all_batch_like:
            missing.add(ref)
            errors.append(f"{sheet}!row{r}: batch_ref {ref} not found in any batch/planting register")
    print(f"{sheet} unresolved batch_ref:", missing or "none")

# ---------------------------------------------------------------------
# 3. Weekly feed_kg sums vs 06_FEED_INVENTORY.used_kg (piggery + poultry)
# ---------------------------------------------------------------------
import datetime as dt

def monday_of(iso_date):
    dd = dt.date.fromisoformat(iso_date)
    return (dd - dt.timedelta(days=dd.weekday())).isoformat()

pig_feed_by_week = defaultdict(float)
for r, vals in sheet_rows("P2_PIG_DAILY_LOG"):
    date, feed_kg = vals[0], vals[4]
    pig_feed_by_week[monday_of(date)] += feed_kg

poultry_feed_by_week = defaultdict(float)
for r, vals in sheet_rows("C2_POULTRY_DAILY_LOG"):
    date, feed_kg = vals[0], vals[7]
    poultry_feed_by_week[monday_of(date)] += feed_kg

tie_checks = []
for r, vals in sheet_rows("06_FEED_INVENTORY"):
    date, domain, used_kg = vals[0], vals[1], vals[5]
    if domain == "piggery":
        actual = pig_feed_by_week.get(date, 0)
    else:
        actual = poultry_feed_by_week.get(date, 0)
    diff = abs(actual - used_kg)
    tie_checks.append((date, domain, used_kg, round(actual, 1), round(diff, 1)))
    if diff > 0.5:
        errors.append(f"06_FEED_INVENTORY week {date} {domain}: used_kg={used_kg} vs summed daily feed_kg={actual:.1f} (diff {diff:.1f})")

print(f"\nFeed tie-out sample (first 6 of {len(tie_checks)}):")
for row in tie_checks[:6]:
    print(" ", row)
bad_ties = [t for t in tie_checks if t[4] > 0.5]
print(f"Feed tie-out mismatches (>0.5kg): {len(bad_ties)} / {len(tie_checks)}")

# ---------------------------------------------------------------------
# 4. Findings confirmation
# ---------------------------------------------------------------------
print("\n--- FINDING 1: poultry bad batch (heatwave mortality) ---")
for r, vals in sheet_rows("C1_POULTRY_BATCHES", start=3):
    print(" batch", vals[0], "placed", vals[4], "off", vals[8], "status", vals[9])
deaths_by_batch = defaultdict(int)
placed_by_batch = {}
for r, vals in sheet_rows("C1_POULTRY_BATCHES", start=3):
    placed_by_batch[vals[0]] = vals[5]
for r, vals in sheet_rows("C2_POULTRY_DAILY_LOG"):
    deaths_by_batch[vals[1]] += vals[3] or 0
for bc, deaths in deaths_by_batch.items():
    placed = placed_by_batch.get(bc, 1)
    print(f"  {bc}: total deaths={deaths}, placed={placed}, mortality%={100*deaths/placed:.1f}")

print("\n--- FINDING 2: piggery disease outbreak ---")
for r, vals in sheet_rows("07_HEALTH_LOG"):
    if vals[4] in ("Treatment", "Inspection") and vals[1] == "piggery":
        print(" ", vals[0], vals[2], vals[4], vals[5], "cost", vals[10], "outcome", vals[12])

print("\n--- FINDING 3: unpaid crop debtor ---")
for r, vals in sheet_rows("03_REVENUE"):
    if vals[12] == "Owing":
        print(" ", vals[0], vals[2], vals[10], vals[11], vals[12], "batch_ref", vals[13])

print("\n--- QUIRK: missing daily-log week (PIG-B04) ---")
b04_dates = sorted(vals[0] for r, vals in sheet_rows("P2_PIG_DAILY_LOG") if vals[1] == "PIG-B04")
gaps = []
prev = None
for dstr in b04_dates:
    dd = dt.date.fromisoformat(dstr)
    if prev and (dd - prev).days > 1:
        gaps.append((prev.isoformat(), dd.isoformat(), (dd - prev).days - 1))
    prev = dd
print(" gaps found:", gaps)

print("\n--- QUIRK: feed price spike ---")
for r, vals in sheet_rows("06_FEED_INVENTORY"):
    if vals[1] == "piggery":
        pass
prices = [(vals[0], vals[8]) for r, vals in sheet_rows("06_FEED_INVENTORY") if vals[1] == "piggery"]
avg_price = sum(p[1] for p in prices) / len(prices)
spikes = [p for p in prices if p[1] > avg_price * 1.15]
print(" avg piggery unit_cost_per_kg:", round(avg_price, 3), " spikes:", spikes)

print("\n--- QUIRK: briefly-late-but-paid (contract channel) ---")
for r, vals in sheet_rows("03_REVENUE"):
    if vals[11] == "Contract":
        print(" ", vals[0], vals[2], vals[10], vals[11], vals[12])

# ---------------------------------------------------------------------
print(f"\n\n=== TOTAL ERRORS FOUND: {len(errors)} ===")
for e in errors[:50]:
    print(" -", e)
